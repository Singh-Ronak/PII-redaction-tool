import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from pii_redactor.docx_text import apply_replacements, iter_text_units


def _doc_with_runs(*pieces):
    document = docx.Document()
    p = document.add_paragraph()
    for text, bold in pieces:
        p.add_run(text).bold = bold
    return document, p


def test_replacement_across_runs_keeps_formatting_and_paragraph_count():
    document, p = _doc_with_runs(("Contact ", False), ("Rashi ", True), ("Patil", True), (" today", False))
    unit = next(iter_text_units(document))
    assert unit.text == "Contact Rashi Patil today"
    apply_replacements(unit.segments, [(8, 19, "John Doe")])
    assert p.text == "Contact John Doe today"
    assert [r.bold for r in p.runs] == [False, True, True, False]
    assert len(document.paragraphs) == 1


def test_replacement_spanning_a_tab():
    document = docx.Document()
    p = document.add_paragraph()
    run = p.add_run("ICICI")
    run._r.append(OxmlElement("w:tab"))
    p.add_run("Securities Limited")
    unit = next(iter_text_units(document))
    assert unit.text == "ICICI\tSecurities Limited"
    apply_replacements(unit.segments, [(0, len(unit.text), "Rege Securities Limited")])
    assert p.text == "Rege Securities Limited"
    assert not p._p.findall(".//" + qn("w:tab"))


def test_multiple_replacements_right_to_left():
    document, p = _doc_with_runs(("Email a@x.com or call +91 98765 43210.", False))
    unit = next(iter_text_units(document))
    t = unit.text
    apply_replacements(unit.segments, [
        (t.index("a@x.com"), t.index("a@x.com") + 7, "john.doe@example.com"),
        (t.index("+91"), t.index("+91") + 15, "+91 12345 67890"),
    ])
    assert p.text == "Email john.doe@example.com or call +91 12345 67890."


def test_table_cells_and_hyperlink_runs_are_units():
    document = docx.Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "DIN"
    table.cell(1, 0).text = "00135070"
    p = document.add_paragraph("Mail ")
    link = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "ksh.ipo@nuvama.com"
    r.append(t)
    link.append(r)
    p._p.append(link)
    units = {u.text: u for u in iter_text_units(document)}
    assert units["00135070"].context == ("DIN",)
    assert "Mail ksh.ipo@nuvama.com" in units  # python-docx's Paragraph.text would drop the hyperlink run
