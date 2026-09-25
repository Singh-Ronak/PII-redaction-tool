import io
import json

import docx
import numpy as np
import pytest
from PIL import Image

from pii_redactor.config import RedactionConfig
from pii_redactor.pipeline import DocxRedactor

from evaluation.make_synthetic import DOCX_PATH, GOLD_PATH


@pytest.fixture(scope="module")
def redactor(detector):
    return DocxRedactor(RedactionConfig(redact_images=False), detector)


def _texts(path):
    return [p.text for p in docx.Document(str(path)).paragraphs]


def test_end_to_end_text_is_replaced_in_place(tmp_path, redactor):
    src = tmp_path / "in.docx"
    d = docx.Document()
    d.add_paragraph("Contact Person: Rashi Patil; E-mail: rashi.patil@gmail.com; Telephone: +91 98765 43210")
    d.add_paragraph("RASHI PATIL is the promoter of Brightline Logistics Private Limited.")
    d.add_paragraph("The Company Secretary and the Directors approved the Offer.")
    d.save(src)
    out = tmp_path / "out.docx"
    report = redactor.redact(src, out, tmp_path / "report.json")

    before, after = _texts(src), _texts(out)
    assert len(after) == len(before)  # nothing appended, nothing removed
    joined = " ".join(after)
    for leaked in ("Rashi", "RASHI", "Patil", "rashi.patil@gmail.com", "98765 43210", "Brightline"):
        assert leaked not in joined
    assert after[2] == before[2]  # role titles are not PII
    fake_name = after[0].split("Contact Person: ")[1].split(";")[0]
    assert after[1].startswith(fake_name.upper())  # same person, same fake, original letter case
    assert "Private Limited" in after[1]
    assert report["text_redactions_by_type"]["PERSON"] == 2
    assert "Rashi" not in (tmp_path / "report.json").read_text()


def test_mapping_file_is_opt_in(tmp_path, detector):
    src = tmp_path / "in.docx"
    d = docx.Document()
    d.add_paragraph("Contact Person: Rashi Patil")
    d.save(src)
    DocxRedactor(RedactionConfig(redact_images=False), detector).redact(src, tmp_path / "a.docx")
    assert not list(tmp_path.glob("*mapping*"))
    cfg = RedactionConfig(redact_images=False, mapping_path=tmp_path / "mapping.json")
    DocxRedactor(cfg, detector).redact(src, tmp_path / "b.docx")
    assert "rashi patil" in json.loads((tmp_path / "mapping.json").read_text())["PERSON"]


def test_bad_input_is_rejected(tmp_path, redactor):
    with pytest.raises(FileNotFoundError):
        redactor.redact(tmp_path / "missing.docx", tmp_path / "out.docx")
    bad = tmp_path / "notes.txt"
    bad.write_text("hello")
    with pytest.raises(ValueError):
        redactor.redact(bad, tmp_path / "out.docx")


def test_id_card_image_blob_is_overwritten(tmp_path, detector):
    pytest.importorskip("pytesseract")
    import pytesseract

    try:
        pytesseract.get_tesseract_version()
    except Exception:
        pytest.skip("Tesseract not installed")
    out = tmp_path / "out.docx"
    DocxRedactor(RedactionConfig(), detector).redact(DOCX_PATH, out)
    gold = json.loads(GOLD_PATH.read_text(encoding="utf-8"))["image"]
    parts = [p for p in docx.Document(str(out)).part.package.iter_parts() if p.content_type.startswith("image/")]
    img = np.asarray(Image.open(io.BytesIO(parts[0].blob)).convert("RGB"))
    assert (img.shape[1], img.shape[0]) == tuple(gold["size"])
    black = np.all(img < 40, axis=2)
    for box in gold["boxes"]:
        l, t, r, b = (int(round(v)) for v in box["box"])
        coverage = black[t:b, l:r].mean()
        assert (coverage >= 0.9) if box["pii"] else (coverage < 0.5), box["text"]
